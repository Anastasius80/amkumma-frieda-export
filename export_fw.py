#!/usr/bin/env python3
"""export_fw.py – FriedaWeb-Datenbank aus MS-SQL-Backup wiederherstellen und nach SQLite und Excel exportieren

Aufruf
------
Normalerweise nicht direkt, sondern über export_fw.sh. Das Shell-Skript startet einen temporären
SQL-Server-Container, ruft dieses Skript auf und entfernt den Container danach wieder.

Direkter Aufruf, wenn bereits ein SQL Server läuft, der die Sicherung unter /backup/<dateiname> sieht:

    MSSQL_HOST=<host> MSSQL_SA_PASSWORD=<passwort> python3 export_fw.py <backup.bak> <ausgabeordner>

    MSSQL_HOST          Adresse des SQL Servers (Port 1433)
    MSSQL_SA_PASSWORD   Passwort des Benutzers sa
    <backup.bak>        Pfad zur Sicherung auf diesem Rechner; für den SQL Server zählt nur der
                        Dateiname, er muss die Datei unter /backup/<dateiname> finden
    <ausgabeordner>     muss existieren

Benötigte Python-Pakete: pymssql, openpyxl (export_fw.sh installiert sie in .venv).

Ablauf
------
1. Wiederherstellen
   Das Skript wartet bis zu drei Minuten, bis der SQL Server Verbindungen annimmt. Den
   Datenbanknamen liest es aus der Sicherung (RESTORE HEADERONLY), die enthaltenen Daten- und
   Protokolldateien ebenfalls (RESTORE FILELISTONLY). Diese werden beim RESTORE nach
   /var/opt/mssql/data umgeleitet, weil die Originalpfade (C:\\Program Files\\...) unter Linux nicht
   existieren. Eine ältere Sicherung (z. B. SQL Server 2017) wird dabei automatisch auf das Format
   des laufenden Servers aktualisiert; die Daten bleiben unverändert.

2. SQLite – vollständige Kopie
   Jede Tabelle der Datenbank wird mit allen Zeilen übernommen, auch leere Tabellen. Spaltennamen,
   Primärschlüssel und Fremdschlüssel entsprechen dem Original. Typen werden so abgebildet:
       int, smallint, tinyint, bit   -> INTEGER (bit als 0/1)
       float                         -> REAL
       money                         -> NUMERIC (als Gleitkommazahl gespeichert)
       varbinary                     -> BLOB
       datetime                      -> TEXT im Format "YYYY-MM-DD HH:MM:SS.mmm"
       nvarchar, varchar, char       -> TEXT
   Beispiel für eine Abfrage:
       sqlite3 FW_Beispiel_20260101.sqlite \\
         "SELECT g.GrabNr, v.Nachname, v.Vorname, v.Sterbedatum
            FROM tbl_Verstorbener v JOIN tbl_Grab g ON g.ID = v.tbl_Grab_ID
           WHERE v.Nachname LIKE 'Mayer%'"

3. Excel – Arbeitsdatei
   Blatt "Info": Quelle, Friedhof, letzter Sterbefall, letzte Änderung, Zeilenzahl je Blatt.
   Danach je ein Blatt pro Eintrag in SHEETS:
       Gräber               tbl_Grab                 Grabstellen mit Grabart, Reihe, Nutzungsdauer
       Verstorbene          tbl_Verstorbener         Sterbefälle und Bestattungen
       Verstorbene_Archiv   tbl_VerstorbenerArchiv   frühere Belegungen aufgelassener Gräber
       Berechtigte          tbl_Berechtigter         Nutzungsberechtigte mit Adresse und Bankdaten
       Berechtigte_Archiv   tbl_BerechtigterArchiv   frühere Nutzungsberechtigte
       Nutzungsverlauf      tbl_GrabNutzVerlauf      Beginn, Ende und Verlängerungen der Nutzung
       Dokumente            tbl_Bild                 zugeordnete Dokumente (nur Titel und Pfad)
       Verlauf              tbl_Verlauf              Zugriffsprotokoll der Anwendung
       Ereignisprotokoll    tbl_EventProtokoll       Änderungsprotokoll der Anwendung
   Aufbereitung:
     - Spalten mit Fremdschlüssel auf eine Nachschlagetabelle enthalten deren Klartext statt der ID;
       die Überschrift verliert dabei Präfix und "_ID" (tblxx_Konfession_ID -> Konfession). Als
       Klartext dient die erste befüllte Spalte aus LABEL_COLS (Bezeichnung, Firmenname, Firma, Block).
     - Verweise auf ein Grab behalten die numerische ID (tbl_Grab_ID bzw. Grab_ID; im Blatt
       Gräber die Spalte ID) und bekommen daneben die Spalte "Grab" im Format "Reihe / Grabnummer",
       z. B. "Wand 5C / 5". Über die ID lassen sich die Blätter per SVERWEIS/XVERWEIS verbinden.
     - Im Blatt Dokumente ist das Originalfeld "Titel#Pfad" in die Spalten Titel und Pfad geteilt.
     - Datumswerte sind echte Excel-Datumswerte (TT.MM.JJJJ, mit Uhrzeit nur wenn vorhanden).
       Datumswerte vor 1900 kann Excel nicht darstellen; sie stehen als Text in der Zelle.
     - Steuerzeichen, die Excel nicht speichern kann, werden aus Texten entfernt.
     - Erste Zeile fixiert, Autofilter auf allen Spalten.
   Ein weiteres Blatt lässt sich ergänzen, indem man ein Paar (Blattname, Tabellenname) in SHEETS
   einträgt; die Aufbereitung geschieht automatisch anhand der Fremdschlüssel.

4. Prüfungen
   Abbruch mit AssertionError, wenn
     - die Zeilenzahl einer Tabelle in SQLite von der Quelle abweicht,
     - SQLite Fremdschlüssel-Verletzungen meldet (PRAGMA foreign_key_check),
     - ein Text länger als die 32.767 Zeichen ist, die eine Excel-Zelle aufnehmen kann.
   Verweise, zu denen es keinen Nachschlagewert gibt, erscheinen in Excel als "?<ID>"; ihre
   Anzahl wird am Ende als WARNUNG ausgegeben.

Bekannte Eigenheiten der Daten
------------------------------
- 1900-01-01 wurde in FriedaWeb als Platzhalter für unbekannte Daten verwendet und ist
  unverändert übernommen.
- Die Pfade im Blatt Dokumente zeigen auf D:\\FriedaWeb-Projekte\\... des ehemaligen Servers;
  die Dateien selbst sind nicht in der Sicherung enthalten.

SPDX-License-Identifier: MIT
Copyright (c) 2026 Alexander Waller

Permission is hereby granted, free of charge, to any person obtaining a copy
of this software and associated documentation files (the "Software"), to deal
in the Software without restriction, including without limitation the rights
to use, copy, modify, merge, publish, distribute, sublicense, and/or sell
copies of the Software, and to permit persons to whom the Software is
furnished to do so, subject to the following conditions:

The above copyright notice and this permission notice shall be included in all
copies or substantial portions of the Software.

THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR
IMPLIED, INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY,
FITNESS FOR A PARTICULAR PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE
AUTHORS OR COPYRIGHT HOLDERS BE LIABLE FOR ANY CLAIM, DAMAGES OR OTHER
LIABILITY, WHETHER IN AN ACTION OF CONTRACT, TORT OR OTHERWISE, ARISING FROM,
OUT OF OR IN CONNECTION WITH THE SOFTWARE OR THE USE OR OTHER DEALINGS IN THE
SOFTWARE.
"""
import datetime as dt
import decimal
import os
import re
import sqlite3
import sys
import time

import pymssql
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BAK, OUT = sys.argv[1], sys.argv[2]
HOST, PW = os.environ["MSSQL_HOST"], os.environ["MSSQL_SA_PASSWORD"]
STEM = os.path.splitext(os.path.basename(BAK))[0]

SHEETS = [
    ("Gräber", "tbl_Grab"),
    ("Verstorbene", "tbl_Verstorbener"),
    ("Verstorbene_Archiv", "tbl_VerstorbenerArchiv"),
    ("Berechtigte", "tbl_Berechtigter"),
    ("Berechtigte_Archiv", "tbl_BerechtigterArchiv"),
    ("Nutzungsverlauf", "tbl_GrabNutzVerlauf"),
    ("Dokumente", "tbl_Bild"),
    ("Verlauf", "tbl_Verlauf"),
    ("Ereignisprotokoll", "tbl_EventProtokoll"),
]
LABEL_COLS = ["Bezeichnung", "Firmenname", "Firma", "Block"]


def connect(db="master"):
    for _ in range(90):
        try:
            return pymssql.connect(server=HOST, user="sa", password=PW, database=db, charset="UTF-8", autocommit=True)
        except pymssql.Error:
            time.sleep(2)
    sys.exit("SQL Server nicht erreichbar")


# --- Restore ---
mcur = connect().cursor(as_dict=True)
src = "/backup/" + os.path.basename(BAK)
mcur.execute("RESTORE HEADERONLY FROM DISK=%s", src)
DB = mcur.fetchall()[0]["DatabaseName"]
mcur.execute("RESTORE FILELISTONLY FROM DISK=%s", src)
moves = ", ".join(
    f"MOVE N'{f['LogicalName']}' TO N'/var/opt/mssql/data/{DB}_{i}.{'ldf' if f['Type'] == 'L' else 'mdf'}'"
    for i, f in enumerate(mcur.fetchall()))
print(f"Restore {DB} aus {BAK} …", flush=True)
mcur.execute(f"RESTORE DATABASE [{DB}] FROM DISK=%s WITH {moves}", src)
while True:
    mcur.execute("SELECT state_desc FROM sys.databases WHERE name=%s", DB)
    if (mcur.fetchone() or {}).get("state_desc") == "ONLINE":
        break
    time.sleep(1)

cur = connect(DB).cursor()


def q(sql):
    cur.execute(sql)
    return cur.fetchall()


tables = [r[0] for r in q("SELECT name FROM sys.tables ORDER BY name")]
cols = {}
for t, c, typ in q("SELECT TABLE_NAME, COLUMN_NAME, DATA_TYPE FROM INFORMATION_SCHEMA.COLUMNS ORDER BY TABLE_NAME, ORDINAL_POSITION"):
    cols.setdefault(t, []).append((c, typ))
pks = {}
for t, c in q("""SELECT t.name, c.name FROM sys.indexes i JOIN sys.index_columns ic ON ic.object_id=i.object_id AND ic.index_id=i.index_id
                 JOIN sys.columns c ON c.object_id=ic.object_id AND c.column_id=ic.column_id JOIN sys.tables t ON t.object_id=i.object_id
                 WHERE i.is_primary_key=1 ORDER BY t.name, ic.key_ordinal"""):
    pks.setdefault(t, []).append(c)
fks = {}
for child, col, parent, pcol in q("""SELECT DISTINCT OBJECT_NAME(fkc.parent_object_id), COL_NAME(fkc.parent_object_id, fkc.parent_column_id),
                 OBJECT_NAME(fkc.referenced_object_id), COL_NAME(fkc.referenced_object_id, fkc.referenced_column_id)
                 FROM sys.foreign_key_columns fkc"""):
    if (child, col) != (parent, pcol):
        fks.setdefault(child, {})[col] = (parent, pcol)


def fetch(t):
    cur.execute(f"SELECT * FROM [{t}]")
    return cur.fetchall()


# --- SQLite: every table, original names/values, PKs and FKs declared ---
SQLITE_TYPE = {"int": "INTEGER", "smallint": "INTEGER", "tinyint": "INTEGER", "bit": "INTEGER",
               "float": "REAL", "money": "NUMERIC", "varbinary": "BLOB"}
sqlite3.register_adapter(decimal.Decimal, float)
sqlite3.register_adapter(dt.datetime, lambda d: d.isoformat(" ", "milliseconds"))

db_path = os.path.join(OUT, STEM + ".sqlite")
if os.path.exists(db_path):
    os.remove(db_path)
lite = sqlite3.connect(db_path)
data = {}
for t in tables:
    defs = [f'"{c}" {SQLITE_TYPE.get(typ, "TEXT")}' for c, typ in cols[t]]
    if t in pks:
        defs.append("PRIMARY KEY (" + ", ".join(f'"{c}"' for c in pks[t]) + ")")
    for c, (p, pc) in fks.get(t, {}).items():
        defs.append(f'FOREIGN KEY ("{c}") REFERENCES "{p}" ("{pc}")')
    lite.execute(f'CREATE TABLE "{t}" (\n  ' + ",\n  ".join(defs) + "\n)")
    rows = data[t] = fetch(t)
    lite.executemany(f'INSERT INTO "{t}" VALUES ({",".join("?" * len(cols[t]))})', rows)
lite.commit()
for t in tables:
    n = lite.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0]
    assert n == len(data[t]), (t, n, len(data[t]))
assert not lite.execute("PRAGMA foreign_key_check").fetchall(), "FK-Verletzungen in SQLite"
lite.close()


# --- Excel: core tables, FK ids replaced by readable labels ---
def label_map(t):
    names = [c for c, _ in cols[t]]
    if t == "tbl_Grab":
        reihe = label_map("tblxx_Reihe")
        i_r, i_nr = names.index("tblxx_Reihe_ID"), names.index("GrabNr")
        return {r[0]: f"{reihe.get(r[i_r], '')} / {r[i_nr]}" for r in data[t]}
    idx = [names.index(c) for c in LABEL_COLS if c in names]
    return {r[0]: next((r[i] for i in idx if r[i] not in (None, "")), r[0]) for r in data[t]}


def header_for(col):
    return re.sub(r"(?i)^(tblxx_|loktbl_|tbl_)|_ID$", "", col)


def cell(ws, v, bold=False):
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
        assert len(v) <= 32767, "text too long for Excel"
    elif isinstance(v, dt.datetime) and v.year < 1900:
        v = v.isoformat(" ")
    c = WriteOnlyCell(ws, v)
    if isinstance(v, dt.datetime):
        c.number_format = "DD.MM.YYYY" if v.time() == dt.time() else "DD.MM.YYYY HH:MM:SS"
    if bold:
        c.font = Font(bold=True)
    return c


wb = Workbook(write_only=True)
info = wb.create_sheet("Info")
sheet_counts = []
dangling = 0

for name, t in SHEETS:
    names = [c for c, _ in cols[t]]
    maps = {c: label_map(p) for c, (p, _) in fks.get(t, {}).items()}
    if t == "tbl_Grab":
        maps["ID"] = label_map("tbl_Grab")
    out_cols = []  # (header, src index, label map / function / None)
    for i, c in enumerate(names):
        p = fks.get(t, {}).get(c, (None,))[0]
        if p == "tbl_Grab" or (t == "tbl_Grab" and c == "ID"):
            out_cols += [(c, i, None), ("Grab" if c in ("ID", "tbl_Grab_ID", "Grab_ID") else header_for(c), i, maps[c])]
        elif c in maps:
            out_cols.append((header_for(c), i, maps[c]))
        elif t == "tbl_Bild" and c == "Bildpfad":
            out_cols += [("Titel", i, lambda v: v.split("#", 1)[0]), ("Pfad", i, lambda v: v.split("#", 1)[-1])]
        else:
            out_cols.append((c, i, None))

    ws = wb.create_sheet(name)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(out_cols))}{len(data[t]) + 1}"
    ws.append([cell(ws, h, bold=True) for h, _, _ in out_cols])
    for r in data[t]:
        row = []
        for _, i, m in out_cols:
            v = r[i]
            if m is not None and v is not None:
                if callable(m):
                    v = m(v)
                elif v in m:
                    v = m[v]
                else:
                    v, dangling = f"?{v}", dangling + 1
            row.append(cell(ws, v))
        ws.append(row)
    sheet_counts.append((name, t, len(data[t])))

friedhof = ", ".join(f"{b} {o or ''}".strip() for b, o in q("SELECT Bezeichnung, Ort FROM tblxx_Friedhof"))
last_death = q("SELECT MAX(Sterbedatum) FROM tbl_Verstorbener")[0][0]
last_change = q("SELECT MAX(Datum) FROM tbl_Verlauf")[0][0]
info.append([cell(info, f"Export der Friedhofsverwaltung FriedaWeb – {friedhof}", bold=True)])
info.append(["Quelle", f"MS-SQL-Backup {os.path.basename(BAK)} (Datenbank {DB})"])
info.append(["Letzter Sterbefall", last_death])
info.append(["Letzte Änderung", last_change])
info.append(["Vollständige Kopie", f"{STEM}.sqlite enthält alle {len(tables)} Tabellen unverändert mit IDs und Fremdschlüsseln"])
info.append(["Hinweis", "Nachschlage-IDs (Grabart, Reihe, Konfession, Firma …) sind hier durch ihre Bezeichnung ersetzt. "
             "Spalte 'Grab' = Reihe / Grabnummer; tbl_Grab_ID bzw. Grab_ID verweist auf die ID im Blatt Gräber."])
info.append([])
info.append([cell(info, "Blatt", bold=True), cell(info, "Quelltabelle", bold=True), cell(info, "Zeilen", bold=True)])
for row in sheet_counts:
    info.append(list(row))

xlsx_path = os.path.join(OUT, STEM + ".xlsx")
wb.save(xlsx_path)
print(f"SQLite: {db_path} ({len(tables)} Tabellen)")
print(f"Excel:  {xlsx_path}")
for n, t, c in sheet_counts:
    print(f"  {n:<20} {c:>6}")
if dangling:
    print(f"WARNUNG: {dangling} Verweise ohne Nachschlagewert (als ?ID markiert)")
